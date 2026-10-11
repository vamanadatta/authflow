import os
from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

def generate_formatted_excel(records):
    """
    Generates a beautifully formatted Master Excel workbook from database records.
    Returns BytesIO object containing the .xlsx file.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master Accounting Sheet"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY_HEADER = "1E293B"     # Slate Navy
    ACCENT_BLUE = "4338CA"     # Indigo
    ZEBRA_FILL = "F8FAFC"      # Light slate
    BORDER_COLOR = "CBD5E1"    # Soft gray
    TOTAL_FILL = "EEF2F6"      # Subtle highlight

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    double_bottom_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='double', color='0F172A')
    )

    # 1. Title Banner
    ws.merge_cells('A1:K1')
    title_cell = ws['A1']
    title_cell.value = "MASTER ACCOUNTING & EXPENSE REGISTER"
    title_cell.font = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 42

    # 2. Metadata Subheader
    ws.merge_cells('A2:K2')
    sub_cell = ws['A2']
    sub_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Entries: {len(records)} | Formatted Master Sheet"
    sub_cell.font = Font(name='Segoe UI', size=10, italic=True, color='64748B')
    sub_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 22

    # Empty spacer row 3
    ws.row_dimensions[3].height = 12

    # 3. Table Column Headers (Row 4)
    headers = [
        "Entry ID",
        "Date",
        "Vendor / Client",
        "Invoice / Ref #",
        "Description / Particulars",
        "Category",
        "Subtotal ($)",
        "Tax / VAT ($)",
        "Total Amount ($)",
        "Payment Status",
        "Raw Source File"
    ]

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.row_dimensions[4].height = 30
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # 4. Populate Data Rows
    start_row = 5
    current_row = start_row

    for rec in records:
        ws.row_dimensions[current_row].height = 24
        
        # Determine zebra striping
        is_even = (current_row % 2 == 0)
        row_fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type='solid') if is_even else None

        row_data = [
            (rec["id"], "center", None),
            (rec["entry_date"], "center", "YYYY-MM-DD"),
            (rec["vendor_client"], "left", None),
            (rec["invoice_no"] or "-", "center", None),
            (rec["description"], "left", None),
            (rec["category"], "center", None),
            (float(rec["subtotal"]), "right", "$#,##0.00"),
            (float(rec["tax_amount"]), "right", "$#,##0.00"),
            (float(rec["total_amount"]), "right", "$#,##0.00"),
            (rec["payment_status"], "center", None),
            (rec["original_filename"] or "Direct Input", "left", None)
        ]

        for col_idx, (val, align_h, num_fmt) in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = Font(name='Segoe UI', size=10)
            cell.alignment = Alignment(horizontal=align_h, vertical='center')
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if num_fmt:
                cell.number_format = num_fmt

        current_row += 1

    # 5. Total / Summary Row (if records exist)
    if records:
        total_row = current_row
        ws.row_dimensions[total_row].height = 28

        # Merge A to F for 'TOTAL' label
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=6)
        label_cell = ws.cell(row=total_row, column=1, value="TOTAL SUMMARY:")
        label_cell.font = Font(name='Segoe UI', size=11, bold=True, color='0F172A')
        label_cell.alignment = Alignment(horizontal='right', vertical='center')

        # Fill merged cells
        for c in range(1, 7):
            c_cell = ws.cell(row=total_row, column=c)
            c_cell.fill = PatternFill(start_color=TOTAL_FILL, end_color=TOTAL_FILL, fill_type='solid')
            c_cell.border = double_bottom_border

        # Formulas for Subtotal, Tax, Total
        end_data_row = current_row - 1
        formula_cols = [
            (7, f"=SUM(G{start_row}:G{end_data_row})"),
            (8, f"=SUM(H{start_row}:H{end_data_row})"),
            (9, f"=SUM(I{start_row}:I{end_data_row})"),
        ]

        for col_idx, formula in formula_cols:
            f_cell = ws.cell(row=total_row, column=col_idx, value=formula)
            f_cell.font = Font(name='Segoe UI', size=11, bold=True, color='0F172A')
            f_cell.fill = PatternFill(start_color=TOTAL_FILL, end_color=TOTAL_FILL, fill_type='solid')
            f_cell.alignment = Alignment(horizontal='right', vertical='center')
            f_cell.border = double_bottom_border
            f_cell.number_format = "$#,##0.00"

        # Blank for remaining columns
        for c in range(10, 12):
            b_cell = ws.cell(row=total_row, column=c, value="")
            b_cell.fill = PatternFill(start_color=TOTAL_FILL, end_color=TOTAL_FILL, fill_type='solid')
            b_cell.border = double_bottom_border

    # 6. Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip merged banner row
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def parse_raw_file_preview(filepath):
    """
    Reads an uploaded raw Excel/CSV file and returns preview data (headers and top rows).
    """
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext in ['.xlsx', '.xls']:
            df = pd.read_excel(filepath, nrows=10)
        elif ext == '.csv':
            df = pd.read_csv(filepath, nrows=10)
        else:
            return None
        
        # Replace NaN with empty string
        df = df.fillna("")
        return {
            "columns": list(df.columns),
            "rows": df.values.tolist(),
            "total_preview_rows": len(df)
        }
    except Exception as e:
        return {"error": str(e)}

def generate_companies_excel(companies):
    """
    Generates an Excel register of all new company registration applications.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Company Applications"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette
    EMERALD_HEADER = "065F46"
    ACCENT_TEAL = "0D9488"
    BORDER_COLOR = "CBD5E1"
    ZEBRA_FILL = "F8FAFC"

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    # 1. Title Banner
    ws.merge_cells('A1:J1')
    title_cell = ws['A1']
    title_cell.value = "NEW COMPANY REGISTRATION REGISTER"
    title_cell.font = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
    title_cell.fill = PatternFill(start_color=EMERALD_HEADER, end_color=EMERALD_HEADER, fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 42

    # 2. Subheader
    ws.merge_cells('A2:J2')
    sub_cell = ws['A2']
    sub_cell.value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Total Applications: {len(companies)}"
    sub_cell.font = Font(name='Segoe UI', size=10, italic=True, color='64748B')
    sub_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 22

    # Spacer
    ws.row_dimensions[3].height = 12

    # 3. Headers
    headers = [
        "App ID",
        "Company Name",
        "Business Structure",
        "Industry",
        "Contact Person",
        "Contact Email",
        "Contact Phone",
        "Documents Attached",
        "Application Status",
        "Submission Date"
    ]

    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color=ACCENT_TEAL, end_color=ACCENT_TEAL, fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.row_dimensions[4].height = 30
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # 4. Data
    for row_idx, comp in enumerate(companies, start=5):
        ws.row_dimensions[row_idx].height = 24
        is_even = (row_idx % 2 == 0)
        row_fill = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type='solid') if is_even else None

        row_data = [
            (comp["id"], "center"),
            (comp["company_name"], "left"),
            (comp["business_type"], "center"),
            (comp["industry"] or "-", "center"),
            (comp["contact_person"] or comp["applicant_name"], "left"),
            (comp["contact_email"] or comp["applicant_email"], "left"),
            (comp["contact_phone"] or "-", "center"),
            (f"{comp['doc_count']} Forms Uploaded", "center"),
            (comp["status"], "center"),
            (comp["created_at"][:10] if comp["created_at"] else "-", "center")
        ]

        for col_idx, (val, align_h) in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Segoe UI', size=10)
            cell.alignment = Alignment(horizontal=align_h, vertical='center')
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

